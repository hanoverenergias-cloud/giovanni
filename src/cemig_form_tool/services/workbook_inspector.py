"""Workbook inspection utilities focused on safe-write decisions."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any


@dataclass
class SheetInspection:
    sheet_name: str
    hidden: bool
    formula_cells: list[str]
    locked_cells: list[str]
    editable_candidates: list[str]


class WorkbookInspector:
    def inspect(self, workbook_path: Path, target_sheet: str = "Formulário") -> dict:
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError as exc:  # pragma: no cover - dependency guard
            raise RuntimeError("openpyxl is required for workbook inspection") from exc

        wb = load_workbook(workbook_path, data_only=False)
        report: dict[str, Any] = {
            "sheet_names": wb.sheetnames,
            "hidden_sheets": [],
            "sheets": {},
            "warnings": [],
        }

        for name in wb.sheetnames:
            ws = wb[name]
            hidden = ws.sheet_state != "visible"
            if hidden:
                report["hidden_sheets"].append(name)
            report["sheets"][name] = self._inspect_sheet(ws)

        if target_sheet not in wb.sheetnames:
            report["warnings"].append(f"Target sheet '{target_sheet}' not found")

        return report

    def _inspect_sheet(self, ws: Any) -> SheetInspection:
        formulas: list[str] = []
        locked: list[str] = []
        editable: list[str] = []
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append(cell.coordinate)
                    continue

                if cell.protection and cell.protection.locked:
                    locked.append(cell.coordinate)
                elif cell.value in (None, ""):
                    editable.append(cell.coordinate)

        return SheetInspection(
            sheet_name=ws.title,
            hidden=ws.sheet_state != "visible",
            formula_cells=formulas,
            locked_cells=locked,
            editable_candidates=editable,
        )
