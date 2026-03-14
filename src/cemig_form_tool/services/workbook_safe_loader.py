"""Safe workbook loader that keeps formulas/styles/validations intact."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook


class WorkbookSafeLoader:
    def load_template(self, path: Path) -> Workbook:
        return load_workbook(path, data_only=False, keep_vba=True)

    def clone_template(self, template_path: Path, output_path: Path) -> Workbook:
        wb = self.load_template(template_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        return load_workbook(output_path, data_only=False, keep_vba=True)
