from __future__ import annotations

from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection
from openpyxl.worksheet.datavalidation import DataValidation

from cemig_form_tool.models.schemas import ClientData, FieldMapping
from cemig_form_tool.services.workbook_filler import WorkbookFiller


def build_template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Formulário"
    wb.create_sheet("Dados")
    wb["Dados"].sheet_state = "hidden"

    ws["B6"] = ""
    ws["C6"] = "=SUM(1,2)"
    ws["D6"] = "option"
    ws["E6"] = "locked"
    ws["E6"].protection = Protection(locked=True)

    ws.merge_cells("B9:C9")
    ws["B9"] = ""

    dv = DataValidation(type="list", formula1='"A,B,C"')
    ws.add_data_validation(dv)
    dv.add("D6")

    wb.save(path)


def test_preserves_formulas_hidden_sheets_validations_and_merged_cells(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    output = tmp_path / "output.xlsx"
    build_template(template)

    wb = load_workbook(template)
    data = ClientData(full_name="Maria da Silva", address="Rua 1")
    mappings = [
        FieldMapping(field_name="full_name", sheet="Formulário", cell="B6", required=True),
        FieldMapping(field_name="address", sheet="Formulário", cell="B9"),
        FieldMapping(field_name="city", sheet="Formulário", cell="E6"),
    ]

    writes = WorkbookFiller().fill(wb, data, mappings)
    wb.save(output)
    reloaded = load_workbook(output)
    ws = reloaded["Formulário"]

    assert ws["C6"].value == "=SUM(1,2)"
    assert reloaded["Dados"].sheet_state == "hidden"
    assert len(ws.data_validations.dataValidation) == 1
    assert "B9:C9" in {str(rng) for rng in ws.merged_cells.ranges}
    assert ws["E6"].value == "locked"
    assert len(writes) == 2


def test_never_overwrites_formula_cell_without_explicit_override(tmp_path: Path) -> None:
    template = tmp_path / "template_formula.xlsx"
    build_template(template)
    wb = load_workbook(template)
    data = ClientData(full_name="Cliente")
    mappings = [FieldMapping(field_name="full_name", sheet="Formulário", cell="C6")]

    WorkbookFiller().fill(wb, data, mappings)

    assert wb["Formulário"]["C6"].value == "=SUM(1,2)"
